from datetime import datetime
from docx import Document

OUTPUT_FILE = "RUBRICA_EVIDENCIAS.docx"


def add_title_page(doc: Document) -> None:
    doc.add_heading("Rubrica de Evidencias y Documentacion del Proyecto", level=0)
    doc.add_paragraph("Proyecto: MC Infrastructure Monitor")
    doc.add_paragraph("Tipo de entrega: Rubrica academica en formato Word")
    doc.add_paragraph(f"Fecha de generacion: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")


def add_rubric_section(doc: Document) -> None:
    doc.add_heading("1. Rubrica de evaluacion", level=1)
    doc.add_paragraph(
        "Instrucciones: completar la columna 'Resultado obtenido' durante la demostracion y adjuntar capturas de evidencia."
    )

    headers = [
        "#",
        "Criterio evaluado",
        "Evidencia requerida",
        "Comando o accion",
        "Resultado obtenido",
        "Puntaje max",
    ]

    rows = [
        ["1", "Separacion en 2 servicios", "Existen dos contenedores en ejecucion", "docker compose ps", "", "10"],
        ["2", "API REST funcional", "Respuesta JSON valida", "curl http://localhost:8000/ (o Swagger)", "", "10"],
        ["3", "Dashboard Web funcional", "Vista web cargada", "Abrir http://localhost:8080", "", "10"],
        ["4", "Red Bridge personalizada", "Ambos servicios en monitor-net", "docker network inspect monitor-net", "", "15"],
        ["5", "Resolucion DNS entre contenedores", "Dashboard consulta api-server", "docker compose exec web-server python -c \"import requests; print(requests.get('http://api-server:5000/api/system', timeout=5).status_code)\"", "", "15"],
        ["6", "Administracion de recursos", "Limites CPU/memoria visibles", "docker compose config", "", "10"],
        ["7", "Healthchecks y disponibilidad", "Servicios healthy", "docker compose ps", "", "10"],
        ["8", "Persistencia historica", "Registros en metrics_history", "docker compose exec api-server python -c \"import sqlite3; c=sqlite3.connect('/app/data/metrics.db'); print(c.execute('select count(*) from metrics_history').fetchone()[0])\"", "", "15"],
        ["9", "Endpoint de historial", "API devuelve historial", "Abrir http://localhost:8000/api/history?limit=10", "", "5"],
    ]

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"

    header_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        header_cells[i].text = header

    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value

    doc.add_paragraph("Puntaje total sugerido: 100 puntos")
    doc.add_paragraph("Aprobacion sugerida: 70 puntos")


def add_evidence_checklist(doc: Document) -> None:
    doc.add_heading("2. Checklist de evidencias", level=1)
    items = [
        "Captura de docker compose ps con ambos contenedores Up/Healthy.",
        "Captura de docker network inspect monitor-net con ambos contenedores.",
        "Captura del dashboard cargado en navegador (http://localhost:8080).",
        "Captura del endpoint /api/history devolviendo datos.",
        "Captura de consulta SQLite con conteo de registros historicos.",
    ]
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_documents_explanation(doc: Document) -> None:
    doc.add_heading("3. Explicacion de cada documento del proyecto", level=1)

    explanations = [
        ("docker-compose.yml", "Orquesta los dos servicios, define la red Bridge personalizada monitor-net, variables de entorno, limites de recursos y volumen persistente api-data para la base de datos."),
        ("Readme.md", "Documento principal del proyecto: arquitectura, instrucciones de despliegue, pruebas funcionales y consideraciones tecnicas."),
        ("RUBRICA_EVIDENCIAS.md", "Version editable en Markdown de la rubrica academica con criterios, evidencias y puntajes."),
        ("api-service/app.py", "API REST con FastAPI. Expone endpoints de estado del servidor y endpoint de historial; tambien inicia captura periodica de metricas."),
        ("api-service/monitor.py", "Modulo de monitoreo del sistema basado en psutil para CPU, memoria, disco, red y procesos."),
        ("api-service/metrics_repository.py", "Capa de persistencia SQLite para guardar y consultar metricas historicas."),
        ("api-service/requirements.txt", "Dependencias de Python del servicio API."),
        ("api-service/Dockerfile", "Imagen de la API: instala dependencias, define healthcheck y comando de arranque."),
        ("dashboard-service/app.py", "Aplicacion Flask del dashboard que renderiza la UI y hace proxy de peticiones hacia la API."),
        ("dashboard-service/templates/index.html", "Interfaz web con metricas en tiempo real, graficas y visualizacion de historial CPU/Memoria."),
        ("dashboard-service/requirements.txt", "Dependencias de Python del dashboard."),
        ("dashboard-service/Dockerfile", "Imagen del dashboard: instala dependencias, define healthcheck y comando de ejecucion."),
        ("exportar_rubrica_word.py", "Script utilitario para generar este archivo Word (.docx) de forma automatica."),
    ]

    for path, description in explanations:
        doc.add_paragraph(path, style="List Number")
        doc.add_paragraph(description)


def main() -> None:
    doc = Document()
    add_title_page(doc)
    add_rubric_section(doc)
    add_evidence_checklist(doc)
    add_documents_explanation(doc)
    doc.save(OUTPUT_FILE)
    print(f"Archivo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
