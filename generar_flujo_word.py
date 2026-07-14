from datetime import datetime
from docx import Document

OUTPUT_FILE = "FLUJO_DATOS_APLICACION.docx"


def add_title(doc: Document) -> None:
    doc.add_heading("Explicacion del Flujo de Datos", level=0)
    doc.add_paragraph("Proyecto: MC Infrastructure Monitor")
    doc.add_paragraph(f"Fecha: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")


def add_architecture(doc: Document) -> None:
    doc.add_heading("1. Arquitectura distribuida", level=1)
    doc.add_paragraph(
        "La aplicacion esta compuesta por dos contenedores en una red Docker Bridge personalizada (monitor-net):"
    )
    doc.add_paragraph("- web-server (Flask): interfaz web y capa proxy.")
    doc.add_paragraph("- api-server (FastAPI): recoleccion de metricas y API REST.")
    doc.add_paragraph(
        "La comunicacion entre contenedores se realiza por DNS interno de Docker usando el nombre del servicio api-server."
    )


def add_main_flow(doc: Document) -> None:
    doc.add_heading("2. Flujo principal de datos (tiempo real)", level=1)

    steps = [
        "El usuario abre el dashboard en http://localhost:8080.",
        "Flask entrega la plantilla index.html al navegador.",
        "JavaScript ejecuta refreshAll() y solicita /api/proxy/status.",
        "El web-server reenvia la solicitud a http://api-server:5000/status.",
        "FastAPI obtiene metricas del sistema (CPU, memoria, disco, red, procesos).",
        "FastAPI retorna JSON de estado al web-server.",
        "El web-server devuelve el JSON al navegador.",
        "El frontend actualiza tarjetas y graficas del dashboard.",
    ]

    for step in steps:
        doc.add_paragraph(step, style="List Number")


def add_history_flow(doc: Document) -> None:
    doc.add_heading("3. Flujo de datos historicos (persistencia)", level=1)

    steps = [
        "Un worker periodico en FastAPI genera snapshots de estado.",
        "Cada snapshot se guarda en SQLite (/app/data/metrics.db) mediante metrics_repository.",
        "Cuando el dashboard solicita /api/proxy/history, Flask consulta /history en la API.",
        "La API lee los registros historicos de SQLite y responde en JSON.",
        "El frontend usa esos datos para dibujar la grafica historica de CPU y memoria.",
    ]

    for step in steps:
        doc.add_paragraph(step, style="List Number")


def add_connectivity_flow(doc: Document) -> None:
    doc.add_heading("4. Flujo de validacion de comunicacion entre contenedores", level=1)

    steps = [
        "El usuario presiona el boton Ver comunicacion en el dashboard.",
        "El frontend solicita /api/proxy/connectivity al web-server.",
        "Flask ejecuta diagnostico: resolucion DNS de api-server y peticion HTTP a /health.",
        "El resultado incluye dns_ok, dns_ip, http_ok, http_status y latencia.",
        "El frontend muestra el diagnostico en el panel de conectividad.",
    ]

    for step in steps:
        doc.add_paragraph(step, style="List Number")


def add_data_summary(doc: Document) -> None:
    doc.add_heading("5. Resumen de datos que circulan", level=1)
    doc.add_paragraph("- Datos de monitoreo: CPU, memoria, disco, red, procesos, timestamp.")
    doc.add_paragraph("- Datos historicos: series temporales persistidas en SQLite.")
    doc.add_paragraph("- Datos de conectividad: DNS interno, estado HTTP y latencia entre contenedores.")


def add_closing(doc: Document) -> None:
    doc.add_heading("6. Conclusion", level=1)
    doc.add_paragraph(
        "El flujo de datos esta desacoplado: el frontend nunca consulta directamente la API interna, "
        "siempre pasa por el web-server como proxy. Esto simplifica seguridad, observabilidad y mantenimiento "
        "de una arquitectura de microservicios en contenedores."
    )


def main() -> None:
    doc = Document()
    add_title(doc)
    add_architecture(doc)
    add_main_flow(doc)
    add_history_flow(doc)
    add_connectivity_flow(doc)
    add_data_summary(doc)
    add_closing(doc)
    doc.save(OUTPUT_FILE)
    print(f"Archivo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
